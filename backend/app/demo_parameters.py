"""Generate fictional parameters without reading business files or a database."""
import argparse
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook


def generate(output: Path) -> None:
    """Exclusively create a workbook; every measurement is invented."""
    output = Path(output)
    if output.suffix.lower() != '.xlsx':
        raise ValueError('Output must have an .xlsx extension')
    wb = Workbook()
    products = wb.active
    products.title = '东北欧产品参数表'
    products.append(('brand', 'product', 'category', 'release_time', 'price_rmb',
                     'price_eur', 'screen', 'chip', 'front_camera', 'rear_camera',
                     'battery', 'wired_charging_w', 'wireless_charging_w',
                     'dimensions', 'ip_rating', 'colors', 'top_selling_point',
                     'official_label', 'official_url'))
    products.append(('DemoBrand', 'Demo One', 'fictional phone', '2026-01',
                     2000, 250, '6.1 inch', 'Demo Chip A\nFictional silicon',
                     '12 MP', '48 MP', 4000, 30, 0, '150 x 70 x 8 mm',
                     'IP54', 'Blue', 'Fictional demo only', 'Demo reference',
                     'https://example.com/demo-one'))
    products.append(('DemoBrand', 'Demo Plus', 'fictional phone', '2026-02',
                     3000, 375, '6.5 inch', 'Demo Chip B', '16 MP', '64 MP',
                     5000, 60, 15, '160 x 75 x 8 mm', 'IP68', 'Silver',
                     'Fictional demo only', 'Demo reference',
                     'https://example.com/demo-plus'))
    chips = wb.create_sheet('芯片天梯图')
    chips.append(('rank', 'chip_name', 'score', 'vendor', 'tier', 'source', 'in_use'))
    # Unsorted input demonstrates score sorting and rank tie breaking.
    chips.append((3, 'Demo Chip A', 100, 'DemoVendor', 'demo', 'fictional', '是'))
    chips.append((2, 'Demo Chip C', 200, 'DemoVendor', 'demo', 'fictional', '否'))
    chips.append((1, 'Demo Chip B', 200, 'DemoVendor', 'demo', 'fictional', '是'))
    note = wb.create_sheet('DEMO_NOTICE')
    note.append(('All names, measurements, prices and scores are fictional.',))
    note.append(('Parser demonstration only; not purchasing or ranking evidence.',))
    try:
        buffer = BytesIO()
        wb.save(buffer)
        # Exclusive creation refuses symlinks and race-time collisions too.
        with output.open('xb') as target:
            target.write(buffer.getvalue())
    finally:
        wb.close()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description='Create a fictional parameter workbook')
    parser.add_argument('output', type=Path)
    args = parser.parse_args(argv)
    try:
        generate(args.output)
    except (OSError, ValueError):
        parser.exit(2, 'Cannot create workbook: use a new .xlsx file in an existing writable directory.\n')
    print('Created fictional workbook: 2 products, 3 chips. Existing files are never overwritten.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
